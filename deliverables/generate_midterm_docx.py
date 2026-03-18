from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt


OUTPUT = Path("中期报告.docx")


def set_run_font(run, size=12, bold=False):
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def set_doc_style(doc: Document):
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.font.size = Pt(12)


def add_title(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, size=18, bold=True)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)


def add_paragraph(doc: Document, text: str, indent=True):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    run = p.add_run(text)
    set_run_font(run, size=12, bold=False)


def add_numbered(doc: Document, items):
    for idx, item in enumerate(items, start=1):
        add_paragraph(doc, f"{idx}. {item}")


def main():
    doc = Document()
    set_doc_style(doc)

    add_title(doc, "毕业设计中期报告")

    add_heading(doc, "课题名称")
    add_paragraph(doc, "数据基础设施标识管理系统的设计与实现", indent=False)

    add_heading(doc, "一、主要内容和计划进度")
    add_paragraph(
        doc,
        "本课题围绕“数据基础设施中的统一标识、分布式解析、资源关联与开放接入”展开，"
        "目标是设计并实现一套可部署、可联调、可演示的数据基础设施标识管理系统。"
        "系统既服务于传统的数据资源登记、解析和共享，也为后续大模型、智能体、MCP 工具接入提供标准化底层支撑。"
    )
    add_paragraph(doc, "本课题的主要研究与实现内容包括以下几个方面：", indent=False)
    add_numbered(
        doc,
        [
            "设计统一的数据标识编码规则与校验规则，支撑标识注册、校验、解析、删除、更新等核心能力。",
            "设计数字对象模型，建立标识、元数据、资源实体之间的关联关系，支撑面向数据资源、工具资源、服务资源的统一管理。",
            "实现面向多节点环境的分布式解析与路由机制，支持节点注册、拓扑维护、跨节点批量解析和全局检索。",
            "实现企业级后端服务，提供标识注册、元数据维护、资源绑定、下载地址生成、审计查询等业务接口。",
            "实现认证中心，为系统提供用户注册、登录、令牌签发、公钥发布和用户信息查询等统一认证能力。",
            "实现 MCP 接入层与 DeepResearch 编排能力，使系统能够被智能体或外部研究前端以标准工具方式调用。",
            "完成本地多节点部署、联调测试和答辩演示环境建设。",
        ],
    )
    add_paragraph(doc, "按照开题阶段的总体计划，项目进度分为四个阶段：", indent=False)
    add_numbered(
        doc,
        [
            "需求分析与总体方案设计阶段：完成课题背景分析、系统目标界定、总体架构设计、技术路线选择。",
            "核心模块开发阶段：完成 SDK 内核、DRIRP 协议支撑、企业后端、认证中心等核心模块的实现。",
            "联调与功能扩展阶段：完成多节点联调、MCP 服务接入、DeepResearch 编排与前端展示。",
            "测试与论文整理阶段：完成测试样例、问题修复、文档沉淀、中后期论文撰写和答辩材料整理。",
        ],
    )
    add_paragraph(
        doc,
        "截至目前，课题总体进度与原计划基本一致，已经完成前 3 个阶段的大部分工作，"
        "当前处于“系统完善、问题收口、文档整理和中后期答辩准备”阶段。"
    )
    add_paragraph(doc, "后续计划如下：", indent=False)
    add_numbered(
        doc,
        [
            "继续完善资源绑定链路中的实体校验与异常处理，补齐上传完成阶段的可靠性控制。",
            "完善前端多页面控制台，增强节点拓扑、标识管理、资源绑定、审计查询和研究编排的展示效果。",
            "继续补充多节点场景下的联调测试、手工测试样例和答辩演示脚本。",
            "完成毕业论文第四章系统设计、第五章系统实现、第六章实验测试等内容的整理与撰写。",
        ],
    )

    add_heading(doc, "二、目前已完成情况")
    add_heading(doc, "1. 已完成系统总体架构设计与工程组织", level=2)
    add_paragraph(
        doc,
        "已完成系统整体模块划分与工程重组，当前系统主要由 dris-java-sdk-next、drirp-java-sdk、"
        "dris-enterprise-backend-v2、dris-auth-center、dris-mcp-server、"
        "dris-deepresearch-orchestrator、dris-deepresearch-frontend 等模块组成。"
        "目前各模块已经具备独立运行和组合联调能力，已形成较为完整的工程化系统框架。"
    )

    add_heading(doc, "2. 已完成标识核心与协议核心能力建设", level=2)
    add_numbered(
        doc,
        [
            "完成标识编码规则、格式校验规则和生成逻辑的设计与实现。",
            "完成数字对象、标识元数据、资源模型、枚举类型等核心模型迁移与重构。",
            "完成标识注册、解析、删除、更新等核心能力的抽象与实现。",
            "完成 DRIRP 协议层的消息模型、请求响应结构、客户端与服务端调用基础能力。",
            "完成跨节点通信中的重试、连接复用、可插拔鉴权与防重放扩展设计。",
        ],
    )

    add_heading(doc, "3. 已完成企业后端主要业务功能", level=2)
    add_numbered(
        doc,
        [
            "标识注册、批量注册、合法性校验、存在性校验、单条解析、批量解析、删除。",
            "元数据查询、更新、补丁更新和数字对象查询。",
            "资源绑定、上传会话、资源定位、下载链接生成等资源管理能力。",
            "本地搜索、全局搜索、跨节点解析、节点拓扑与路由维护。",
            "审计日志记录与审计查询接口设计及实现。",
            "多执行器线程池拆分设计，用于隔离解析、搜索、向量同步等不同类型任务。",
        ],
    )

    add_heading(doc, "4. 已完成认证中心、MCP 与 DeepResearch 链路打通", level=2)
    add_numbered(
        doc,
        [
            "认证中心的注册、登录、令牌获取、公钥获取、用户信息查询等功能。",
            "MCP 服务对标识解析、数字对象解析、全局搜索、路由查询、审计查询等能力的工具封装。",
            "DeepResearch 编排服务与前端页面联通，可基于研究任务调用后端和 MCP 能力生成结构化结果。",
        ],
    )

    add_heading(doc, "5. 已完成本地多节点部署与阶段性测试", level=2)
    add_paragraph(doc, "目前已在本地完成 6 节点部署联调，系统已完成阶段性功能测试，当前统计结果如下：")
    add_numbered(
        doc,
        [
            "测试用例共 36 项。",
            "通过用例 35 项。",
            "通过率为 97.22%。",
            "已打通节点启动与注册、标识注册与解析、资源上传与绑定、跨节点解析与全局检索、审计查询、认证访问、MCP 调用与 DeepResearch 前端展示等主链路。",
        ],
    )

    add_heading(doc, "三、存在的主要问题和解决方案")
    add_heading(doc, "1. 资源上传完成阶段的实体校验仍需加强", level=2)
    add_paragraph(
        doc,
        "当前系统已支持上传会话创建、对象存储上传和资源绑定，但在 completeUpload 阶段，"
        "对“对象是否真实存在、文件大小是否一致、摘要是否一致”的校验仍不够严格。"
        "该问题会影响资源共享场景下的数据可信性。"
    )
    add_paragraph(doc, "解决方案：", indent=False)
    add_numbered(
        doc,
        [
            "在上传完成接口中增加对象存在性校验。",
            "增加文件大小一致性校验。",
            "对提供摘要的文件增加哈希一致性校验。",
            "为上传失败、校验失败等场景补充清晰的错误码和测试用例。",
        ],
    )

    add_heading(doc, "2. 前端控制台功能已具雏形，但整体展示仍需进一步优化", level=2)
    add_paragraph(
        doc,
        "当前前端已经能展示研究工作台、节点拓扑、标识注册、资源绑定等主要功能，"
        "但部分页面交互逻辑、信息组织方式、视觉一致性和操作反馈仍有改进空间，"
        "尚未完全达到毕业答辩演示的理想状态。"
    )
    add_paragraph(doc, "解决方案：", indent=False)
    add_numbered(
        doc,
        [
            "继续完善总览、节点拓扑、标识管理、资源绑定、审计中心等页面。",
            "将常用操作改为更清晰的分步引导和结果反馈方式。",
            "优化中文文案、表单项、结果列表和可视化布局，提高演示效果。",
        ],
    )

    add_heading(doc, "3. 多节点鲁棒性与工程化细节仍需补强", level=2)
    add_paragraph(
        doc,
        "虽然当前已经实现多节点解析、相邻节点状态查询、节点注册与基本联通，"
        "但在节点异常下线、重连恢复、批量场景性能观测、连接治理等方面，仍需继续做工程化完善。"
    )
    add_paragraph(doc, "解决方案：", indent=False)
    add_numbered(
        doc,
        [
            "增加更稳定的节点状态探测和异常恢复机制。",
            "继续优化跨节点迭代解析与递归解析两类流程的实现细节。",
            "完善执行器线程池、连接复用、日志与指标采集，降低链路抖动对整体系统的影响。",
        ],
    )

    add_heading(doc, "4. 论文材料和规范化文档仍需继续沉淀", level=2)
    add_paragraph(
        doc,
        "当前开发文档、测试样例、部署说明已经积累较多，但仍偏向工程联调视角，"
        "尚需进一步转化为符合毕业论文写作要求的规范文本。"
    )
    add_paragraph(doc, "解决方案：", indent=False)
    add_numbered(
        doc,
        [
            "继续整理系统设计说明、数据库设计说明、接口设计说明和测试说明。",
            "将现有实现过程中的设计决策、性能优化点和多节点机制转化为论文可用内容。",
            "完成实验设计、指标定义、结果分析和答辩材料统一整理。",
        ],
    )

    add_heading(doc, "四、指导教师意见")
    add_paragraph(
        doc,
        "可填写示例如下：",
        indent=False,
    )
    add_paragraph(
        doc,
        "该生毕业设计进展总体正常，已完成课题总体方案设计、核心模块实现和阶段性联调测试，"
        "系统已形成较为完整的工程框架，能够较好支撑“数据基础设施标识管理系统的设计与实现”这一课题目标。"
        "当前已初步实现标识管理、多节点解析、资源绑定、认证接入、MCP 工具封装及研究前端展示等功能，"
        "阶段成果较为扎实。"
    )
    add_paragraph(
        doc,
        "后续建议学生继续围绕资源校验、前端展示优化、多节点鲁棒性提升和论文材料整理等方面开展工作，"
        "进一步提高系统完整性与论文质量，按计划推进后续测试、论文撰写和答辩准备。"
    )
    add_paragraph(doc, "指导教师签字：____________", indent=False)
    add_paragraph(doc, "日期：____________", indent=False)

    target = OUTPUT
    try:
        doc.save(target)
    except PermissionError:
        target = OUTPUT.with_stem(f"{OUTPUT.stem}-新版")
        doc.save(target)
    print(f"generated: {target.resolve()}")


if __name__ == "__main__":
    main()
