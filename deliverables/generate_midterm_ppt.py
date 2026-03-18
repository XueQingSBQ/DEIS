from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_VERTICAL_ANCHOR
from pptx.util import Inches, Pt


OUTPUT = Path("中期答辩PPT-图示版.pptx")

TITLE = RGBColor(23, 37, 84)
ACCENT = RGBColor(21, 120, 92)
TEXT = RGBColor(45, 45, 45)
MUTED = RGBColor(112, 112, 112)
BG = RGBColor(246, 248, 251)
WHITE = RGBColor(255, 255, 255)
BORDER = RGBColor(215, 223, 232)
LIGHT = RGBColor(233, 241, 248)
GREEN_LIGHT = RGBColor(232, 245, 238)
NAVY_LIGHT = RGBColor(236, 241, 250)
GOLD = RGBColor(201, 157, 61)
SHADOW = RGBColor(226, 232, 240)


def add_bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_title(slide, title, subtitle=""):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.3), Inches(8.8), Inches(0.7))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(24)
    r.font.bold = True
    r.font.color.rgb = TITLE
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.58), Inches(0.9), Inches(8.4), Inches(0.35))
        p2 = sub.text_frame.paragraphs[0]
        r2 = p2.add_run()
        r2.text = subtitle
        r2.font.size = Pt(11)
        r2.font.color.rgb = MUTED


def add_footer(slide, text):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.02), Inches(8.8), Inches(0.22))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = text
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED


def add_textbox(slide, left, top, width, height, text, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_VERTICAL_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box


def add_bullets(slide, items, left=0.8, top=1.55, width=8.3, height=4.8, font_size=18):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = TEXT
        p.space_after = Pt(8)
    return box


def add_card(slide, left, top, width, height, title, body, fill=WHITE, title_color=TITLE, body_color=TEXT):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = BORDER
    tf = shape.text_frame
    tf.clear()
    if title:
        p1 = tf.paragraphs[0]
        p1.alignment = PP_ALIGN.CENTER
        r1 = p1.add_run()
        r1.text = title
        r1.font.size = Pt(12)
        r1.font.bold = True
        r1.font.color.rgb = title_color
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = body
        r2.font.size = Pt(15)
        r2.font.color.rgb = body_color
    else:
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = body
        r.font.size = Pt(15)
        r.font.color.rgb = body_color
    return shape


def add_panel(slide, left, top, width, height, fill=WHITE, border=BORDER, radius_shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shadow = slide.shapes.add_shape(radius_shape, Inches(left + 0.05), Inches(top + 0.06), Inches(width), Inches(height))
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = SHADOW
    shadow.line.color.rgb = SHADOW
    shape = slide.shapes.add_shape(radius_shape, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = border
    return shape


def add_tag(slide, left, top, width, height, text, fill=TITLE, color=WHITE, size=10):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color
    return shape


def add_arrow(slide, x1, y1, x2, y2, color=ACCENT, width=2.2):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    line.line.end_arrowhead = True
    return line


def add_line(slide, x1, y1, x2, y2, color=BORDER, width=1.5):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def build_cover(slide):
    add_bg(slide)
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(1.0), Inches(8.7), Inches(5.0))
    panel.fill.solid()
    panel.fill.fore_color.rgb = WHITE
    panel.line.color.rgb = BORDER
    add_textbox(slide, 1.0, 1.6, 7.8, 0.8, "毕业设计中期答辩", size=28, color=TITLE, bold=True)
    add_textbox(slide, 1.0, 2.45, 7.8, 0.7, "数据基础设施标识管理系统的设计与实现", size=20, color=TEXT)
    add_textbox(slide, 1.0, 4.15, 5.8, 1.0, "学生：__________    学号：__________\n专业：__________    日期：2026 年 3 月", size=16, color=MUTED)
    add_footer(slide, "Midterm Defense")


def build_problem(slide):
    add_bg(slide)
    add_title(slide, "1. 课题背景与问题提出", "先回答为什么做，而不是先介绍模块")
    add_card(slide, 0.8, 1.55, 1.9, 1.3, "问题一", "数据资源分散\n缺少统一标识", fill=LIGHT)
    add_card(slide, 2.95, 1.55, 1.9, 1.3, "问题二", "跨节点环境下\n定位与解析困难", fill=LIGHT)
    add_card(slide, 5.1, 1.55, 1.9, 1.3, "问题三", "共享链路不清晰\n资源关联弱", fill=LIGHT)
    add_card(slide, 7.25, 1.55, 1.9, 1.3, "问题四", "智能体接入缺少\n统一开放接口", fill=LIGHT)
    add_bullets(slide, [
        "现有数据资源通常分散在不同机构、不同节点、不同系统中，名称、路径和管理方式不统一。",
        "当资源跨节点流转时，往往只能知道“有数据”，却难以稳定定位、解析和复用。",
        "因此，本课题聚焦“统一标识 + 数字对象 + 多节点解析 + 开放接入”这一条主线，构建数据基础设施标识管理系统。",
    ], left=0.9, top=3.35, width=8.1, height=2.6, font_size=18)
    add_textbox(slide, 0.9, 6.3, 8.1, 0.35, "本页结论：本课题解决的是多节点数据资源“难标识、难解析、难共享、难开放接入”的问题。", size=12, color=ACCENT, bold=True)
    add_footer(slide, "Slide 2")


def build_targets(slide):
    add_bg(slide)
    add_title(slide, "2. 研究目标与主要内容", "把课题范围收敛到四个可落地目标")
    add_card(slide, 0.8, 1.55, 2.0, 1.55, "目标一", "统一标识编码\n与校验机制", fill=WHITE)
    add_card(slide, 3.0, 1.55, 2.0, 1.55, "目标二", "数字对象模型\n与元数据管理", fill=WHITE)
    add_card(slide, 5.2, 1.55, 2.0, 1.55, "目标三", "多节点解析\n与路由协同", fill=WHITE)
    add_card(slide, 7.4, 1.55, 1.7, 1.55, "目标四", "MCP / 智能体\n开放接入", fill=WHITE)
    add_bullets(slide, [
        "设计统一标识规则，支撑注册、校验、解析、删除、更新等核心操作。",
        "构建数字对象模型，建立标识、元数据、资源之间的关联关系。",
        "实现多节点环境下的路由、转发与跨节点解析能力。",
        "通过认证中心、MCP 和 DeepResearch 前端形成可演示的完整系统链路。",
    ], left=1.0, top=3.75, width=7.9, height=2.0, font_size=18)
    add_textbox(slide, 1.0, 6.28, 8.0, 0.35, "本页结论：本课题不是做一个普通后台，而是做一套面向数据基础设施场景的标识管理能力底座。", size=12, color=ACCENT, bold=True)
    add_footer(slide, "Slide 3")


def build_architecture(slide):
    add_bg(slide)
    add_title(slide, "3. 系统总体架构图", "用一张图讲清系统分层与模块关系")
    add_panel(slide, 0.65, 1.28, 8.75, 5.45, fill=WHITE)
    add_tag(slide, 0.95, 1.52, 1.15, 0.3, "演示入口", fill=ACCENT)
    add_tag(slide, 0.95, 3.95, 1.15, 0.3, "核心业务", fill=TITLE)
    add_tag(slide, 0.95, 5.3, 1.15, 0.3, "多节点底座", fill=GOLD, color=TITLE)

    add_card(slide, 1.35, 1.65, 7.35, 0.68, "", "应用展示层：DeepResearch Frontend / 节点控制台", fill=NAVY_LIGHT)
    add_card(slide, 1.7, 2.48, 6.65, 0.68, "", "研究编排层：DeepResearch Orchestrator", fill=GREEN_LIGHT)
    add_card(slide, 2.0, 3.28, 6.05, 0.68, "", "开放接入层：DRIS MCP Server", fill=WHITE)

    add_panel(slide, 1.3, 4.08, 4.2, 1.0, fill=LIGHT)
    add_textbox(slide, 1.55, 4.3, 3.7, 0.22, "业务服务层", size=11, color=MUTED, bold=True)
    add_textbox(slide, 1.55, 4.55, 3.7, 0.28, "dris-enterprise-backend-v2", size=18, color=TITLE, bold=True, align=PP_ALIGN.CENTER)

    add_panel(slide, 5.75, 4.08, 2.65, 1.0, fill=WHITE)
    add_textbox(slide, 5.95, 4.3, 2.25, 0.22, "认证与访问控制", size=11, color=MUTED, bold=True)
    add_textbox(slide, 5.95, 4.55, 2.25, 0.28, "dris-auth-center", size=16, color=TITLE, bold=True, align=PP_ALIGN.CENTER)

    add_panel(slide, 1.3, 5.38, 2.5, 0.95, fill=WHITE)
    add_textbox(slide, 1.52, 5.64, 2.05, 0.18, "标识核心", size=11, color=MUTED, bold=True)
    add_textbox(slide, 1.5, 5.9, 2.1, 0.2, "dris-java-sdk-next", size=14, color=TITLE, bold=True, align=PP_ALIGN.CENTER)

    add_panel(slide, 4.08, 5.38, 2.5, 0.95, fill=WHITE)
    add_textbox(slide, 4.3, 5.64, 2.05, 0.18, "跨节点协议", size=11, color=MUTED, bold=True)
    add_textbox(slide, 4.28, 5.9, 2.1, 0.2, "drirp-java-sdk", size=14, color=TITLE, bold=True, align=PP_ALIGN.CENTER)

    add_panel(slide, 6.86, 5.38, 1.65, 0.95, fill=WHITE)
    add_textbox(slide, 6.98, 5.63, 1.4, 0.18, "资源存储", size=11, color=MUTED, bold=True)
    add_textbox(slide, 7.02, 5.9, 1.3, 0.2, "COS / 外部存储", size=12, color=TITLE, bold=True, align=PP_ALIGN.CENTER)

    add_panel(slide, 1.2, 6.15, 2.1, 0.4, fill=GREEN_LIGHT)
    add_panel(slide, 3.95, 6.15, 2.1, 0.4, fill=GREEN_LIGHT)
    add_panel(slide, 6.7, 6.15, 2.1, 0.4, fill=GREEN_LIGHT)
    add_textbox(slide, 1.35, 6.25, 1.8, 0.15, "Node A Backend", size=12, color=TITLE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 4.1, 6.25, 1.8, 0.15, "Node B Backend", size=12, color=TITLE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 6.85, 6.25, 1.8, 0.15, "Node C Backend", size=12, color=TITLE, bold=True, align=PP_ALIGN.CENTER)

    add_arrow(slide, 5.05, 3.96, 5.05, 4.05, color=ACCENT, width=2.6)
    add_arrow(slide, 3.45, 5.12, 2.55, 5.38, color=TITLE, width=2.0)
    add_arrow(slide, 4.95, 5.12, 5.25, 5.38, color=TITLE, width=2.0)
    add_arrow(slide, 6.95, 5.0, 7.65, 5.38, color=TITLE, width=2.0)
    add_arrow(slide, 2.55, 6.14, 2.55, 5.98, color=ACCENT, width=1.8)
    add_arrow(slide, 5.0, 6.14, 5.0, 5.98, color=ACCENT, width=1.8)
    add_arrow(slide, 7.75, 6.14, 7.75, 5.98, color=ACCENT, width=1.8)
    add_textbox(slide, 0.9, 6.95, 8.2, 0.22, "本页结论：系统采用分层解耦设计，下层提供标识与协议能力，中层承载业务，上层支持认证、MCP 和前端接入。", size=11, color=ACCENT, bold=True)
    add_footer(slide, "Slide 4")


def build_class_diagram(slide):
    add_bg(slide)
    add_title(slide, "4. 核心类图", "只画最核心的领域模型，不展示实现层细节")
    add_panel(slide, 0.65, 1.3, 8.75, 5.45, fill=WHITE)
    add_tag(slide, 0.95, 1.54, 1.2, 0.3, "领域模型", fill=ACCENT)
    add_panel(slide, 0.95, 2.0, 2.1, 1.95, fill=WHITE)
    add_tag(slide, 1.15, 2.2, 0.9, 0.24, "入口对象", fill=NAVY_LIGHT, color=TITLE)
    add_textbox(slide, 1.15, 2.5, 1.7, 0.25, "Identifier", size=17, color=TITLE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 1.18, 2.88, 1.65, 0.72, "value\nregionCode\norganizationCode\ncheckCode", size=12, color=TEXT, align=PP_ALIGN.CENTER)

    add_panel(slide, 3.15, 1.75, 3.35, 2.55, fill=LIGHT)
    add_tag(slide, 3.45, 2.0, 1.0, 0.24, "聚合根", fill=TITLE)
    add_textbox(slide, 3.5, 2.34, 2.65, 0.28, "DigitalObject", size=21, color=TITLE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 3.72, 2.82, 2.2, 0.85, "identifier\nmetadata\nresources", size=14, color=TEXT, align=PP_ALIGN.CENTER)

    add_panel(slide, 6.8, 2.0, 1.95, 1.95, fill=WHITE)
    add_tag(slide, 7.0, 2.2, 0.95, 0.24, "描述信息", fill=GREEN_LIGHT, color=TITLE)
    add_textbox(slide, 6.98, 2.5, 1.6, 0.28, "IdentifierMetadata", size=15, color=TITLE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 7.0, 2.88, 1.55, 0.8, "name\ndescription\nowner\nindustry\ndataCategory", size=11, color=TEXT, align=PP_ALIGN.CENTER)

    add_panel(slide, 3.85, 4.65, 1.95, 1.55, fill=WHITE)
    add_tag(slide, 4.08, 4.84, 0.9, 0.24, "资源实体", fill=GOLD, color=TITLE)
    add_textbox(slide, 4.02, 5.14, 1.55, 0.22, "Resource", size=16, color=TITLE, bold=True, align=PP_ALIGN.CENTER)
    add_textbox(slide, 4.05, 5.44, 1.48, 0.5, "resourceType\nstorageType\nlocator\nsize/hash", size=11, color=TEXT, align=PP_ALIGN.CENTER)

    add_arrow(slide, 3.05, 2.95, 3.15, 2.95, color=ACCENT, width=2.4)
    add_arrow(slide, 6.5, 2.95, 6.8, 2.95, color=ACCENT, width=2.4)
    add_arrow(slide, 4.8, 4.3, 4.8, 4.62, color=ACCENT, width=2.4)
    add_tag(slide, 2.0, 2.75, 0.85, 0.22, "resolves to", fill=WHITE, color=MUTED, size=9)
    add_tag(slide, 5.95, 2.75, 0.72, 0.22, "contains", fill=WHITE, color=MUTED, size=9)
    add_tag(slide, 4.98, 4.38, 0.72, 0.22, "contains", fill=WHITE, color=MUTED, size=9)

    add_panel(slide, 6.45, 4.55, 2.35, 1.05, fill=GREEN_LIGHT)
    add_textbox(slide, 6.68, 4.78, 1.9, 0.2, "设计含义", size=11, color=MUTED, bold=True)
    add_textbox(slide, 6.68, 5.03, 1.88, 0.35, "解析标识后，不仅返回元数据，还能进一步关联真实资源或服务入口。", size=11, color=TITLE)

    add_panel(slide, 0.9, 6.0, 8.05, 0.45, fill=NAVY_LIGHT)
    add_textbox(slide, 1.1, 6.1, 7.6, 0.18, "标识是入口，数字对象是聚合根，元数据负责描述，资源负责承载实体访问能力。", size=12, color=TITLE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "Slide 5")


def build_sequence(slide):
    add_bg(slide)
    add_title(slide, "5. 关键时序图：跨节点标识解析", "突出系统从单节点解析向多节点协同解析的扩展能力")
    add_panel(slide, 0.65, 1.25, 8.75, 5.5, fill=WHITE)
    add_tag(slide, 0.95, 1.48, 1.25, 0.3, "关键链路", fill=ACCENT)
    participants = [
        ("前端/调用方", 0.9),
        ("当前节点", 3.0),
        ("路由管理", 5.1),
        ("相邻节点", 7.2),
    ]
    for name, left in participants:
        fill = NAVY_LIGHT if name == "当前节点" else WHITE
        add_panel(slide, left, 1.9, 1.35, 0.58, fill=fill)
        add_textbox(slide, left + 0.08, 2.1, 1.2, 0.14, name, size=12, color=TITLE, bold=True, align=PP_ALIGN.CENTER)
        lane_fill = RGBColor(250, 252, 255) if name != "当前节点" else RGBColor(241, 247, 255)
        lane = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left + 0.48), Inches(2.6), Inches(0.38), Inches(3.45))
        lane.fill.solid()
        lane.fill.fore_color.rgb = lane_fill
        lane.line.color.rgb = RGBColor(220, 227, 236)
        add_line(slide, left + 0.67, 2.55, left + 0.67, 6.15, color=RGBColor(176, 188, 205), width=1.2)

    steps = [
        (1.8, 1.6, 3.0, "1. 发起解析请求"),
        (2.45, 3.0, 3.0, "2. 本地查询"),
        (3.1, 3.0, 5.1, "3. 查询候选下一跳"),
        (3.75, 5.1, 3.0, "4. 返回路由信息"),
        (4.45, 3.0, 7.2, "5. DRIRP 远端解析"),
        (5.1, 7.2, 3.0, "6. 返回解析结果"),
        (5.75, 3.0, 1.6, "7. 汇总结果返回"),
    ]
    for y, x1, x2, text in steps:
        color = ACCENT if x1 != x2 else GOLD
        add_arrow(slide, x1 + 0.67, y + 0.45, x2 + 0.67, y + 0.45, color=color, width=2.2)
        label_width = max(abs(x2 - x1) - 0.45, 1.1)
        add_tag(slide, min(x1, x2) + 0.8, y + 0.22, label_width, 0.22, text, fill=WHITE, color=TEXT, size=9)
    add_panel(slide, 0.9, 6.18, 8.15, 0.38, fill=GREEN_LIGHT)
    add_textbox(slide, 1.05, 6.27, 7.8, 0.16, "本地优先，未命中再跨节点转发，这是系统分布式解析能力的核心体现。", size=12, color=TITLE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "Slide 6")


def build_progress(slide):
    add_bg(slide)
    add_title(slide, "6. 目前已完成情况", "中期阶段重点证明：系统不是停留在设计，而是已经可运行、可联调")
    add_card(slide, 0.8, 1.45, 2.05, 1.45, "模块完成", "SDK / 协议\n后端 / 认证\nMCP / Frontend", fill=WHITE)
    add_card(slide, 3.0, 1.45, 2.05, 1.45, "链路打通", "注册\n解析\n资源绑定\n搜索 / 审计", fill=WHITE)
    add_card(slide, 5.2, 1.45, 2.05, 1.45, "多节点联调", "本地 6 节点\n路由 / 跨节点解析", fill=WHITE)
    add_card(slide, 7.4, 1.45, 1.7, 1.45, "测试结果", "36 项\n35 项通过\n97.22%", fill=GREEN_LIGHT)
    add_bullets(slide, [
        "已完成新一代 DRIS 核心 SDK 与 DRIRP 协议 SDK 重构，保留标识规则、核心模型和协议消息能力。",
        "已完成企业后端主要接口：标识注册、数字对象解析、资源绑定、搜索、审计、节点路由、跨节点批量解析。",
        "已完成认证中心、MCP 服务、DeepResearch 编排器与前端工作台的联通。",
        "已完成本地 6 节点部署联调，系统主链路已具备运行和演示条件。",
    ], left=0.95, top=3.45, width=8.0, height=2.4, font_size=17)
    add_textbox(slide, 0.95, 6.35, 8.0, 0.3, "本页结论：当前系统已经从“方案设计阶段”推进到“可运行、可联调、可演示阶段”。", size=12, color=ACCENT, bold=True)
    add_footer(slide, "Slide 7")


def build_plan(slide):
    add_bg(slide)
    add_title(slide, "7. 当前问题、解决方案与后续计划", "中期答辩要明确：还差什么、怎么补、何时完成")
    add_card(slide, 0.65, 1.55, 2.55, 3.9, "当前问题", "1. 上传完成阶段的实体校验仍需加强\n\n2. 前端展示与交互仍需优化\n\n3. 多节点鲁棒性与实验数据仍需补充", fill=WHITE)
    add_card(slide, 3.45, 1.55, 3.0, 3.9, "对应方案", "1. 增加对象存在性、大小一致性和摘要校验\n\n2. 完善节点拓扑、标识管理、资源绑定、审计等页面\n\n3. 继续优化状态探测、连接复用与测试样例", fill=LIGHT)
    add_card(slide, 6.7, 1.55, 2.55, 3.9, "后续安排", "3 月中下旬：问题收口与页面优化\n\n4 月：实验测试与论文第四、五章\n\n5 月：论文定稿与最终答辩准备", fill=GREEN_LIGHT)
    add_textbox(slide, 0.85, 6.02, 8.1, 0.3, "本页结论：后续工作重点不是再铺新功能，而是围绕可靠性、展示性和论文材料做系统收口。", size=12, color=ACCENT, bold=True)
    add_footer(slide, "Slide 8")


def build_summary(slide):
    add_bg(slide)
    add_title(slide, "8. 中期总结", "最后只回答三件事：解决什么问题、做到什么程度、后续怎么完成")
    add_bullets(slide, [
        "本课题面向多节点数据资源场景，聚焦统一标识、数字对象组织、跨节点解析与开放接入问题。",
        "目前已完成系统总体架构设计和核心模块实现，形成了可运行、可联调、可演示的数据基础设施标识管理系统原型。",
        "后续将重点完善资源可信校验、前端展示、实验评估和论文撰写，按计划推进最终答辩。",
    ], left=0.95, top=1.85, width=7.95, height=2.6, font_size=20)
    add_textbox(slide, 1.0, 5.45, 8.0, 0.7, "谢谢各位老师", size=26, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, "Slide 9")


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    builders = [
        build_cover,
        build_problem,
        build_targets,
        build_architecture,
        build_class_diagram,
        build_sequence,
        build_progress,
        build_plan,
        build_summary,
    ]

    for builder in builders:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        builder(slide)

    target = OUTPUT
    try:
        prs.save(target)
    except PermissionError:
        target = OUTPUT.with_stem(f"{OUTPUT.stem}-新版")
        prs.save(target)
    print(f"generated: {target.resolve()}")


if __name__ == "__main__":
    main()
